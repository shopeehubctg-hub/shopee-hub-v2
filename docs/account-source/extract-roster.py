"""Read-only audited snapshot extraction; source workbooks are never changed."""
from pathlib import Path
import hashlib, json
from datetime import datetime
from zoneinfo import ZoneInfo
import openpyxl

root = Path(__file__).resolve().parents[2]
exports = Path('/var/folders/pm/_n98qfm56yvgp555hfjkmbdh0000gn/T/browser-use/exports')
billing = exports / 'WISDOM FACTORY DEPARTMENT BILLING-0c7fde2a-8c7a-4999-9188-b139b953c4c8.xlsx'
master = exports / 'Shopee Hub Master Directory Framework-46563d62-140e-46d5-a743-50539feeca8d.xlsx'
book = openpyxl.load_workbook(billing, read_only=True, data_only=True)
directory = openpyxl.load_workbook(master, read_only=True, data_only=True)
mapping = {r[0]:r for r in list(directory['Store_Market_Mapping'].iter_rows(values_only=True))[1:] if r[0]}
# Explicit audited candidate crosswalk. These are historical candidates, not verified financial joins.
crosswalk = {
 'MY': {'NOMOQ':['MAP0026'],'DRSMILE':['MAP0009'],'ECOPLUS':['MAP0010'],'FUNFFY':['MAP0102'],
 'ILADY':['MAP0011'],'JEEROUL':[],'M PLUS':['MAP0017'],'KATA MARINE':['MAP0014'],
 'MIZINO CHOCOLATE':[],'SCALE STORY':['MAP0028'],'SCALEGEM':['MAP0027'],'TGC':['MAP0101'],
 'MASTERNEV':['MAP0018'],'KIDSAIE':['MAP0006'],'BEYOUTE':['MAP0005'],'MCS':['MAP0019'],
 'MIZINO PLACENTA':['MAP0021','MAP0022'],'CTG4U MALAYSIA':['MAP0008'],'ZEERO SKINCARE':['MAP0030'],
 'SUPU':[],'AGEPROS':['MAP0004'],'URO 360':['MAP0032'],'LIVACT':['MAP0015','MAP0016'],
 'JOURISH':['MAP0012'],'BUGUCARE':['MAP0007'],'ISOKAE':['MAP0095'],'HUMEAL':[],
 'MOESIE':['MAP0023'],'NINOKO':['MAP0025'],'RECOVIT':['MAP0024'],'MFORMULA':['MAP0020'],
 'SKINDAE':['MAP0029'],'YCT':['MAP0031'],'GOHERB':['MAP0002'],
 'BERLANCO BEAUTY OFFICIAL':['MAP0065'],'BONLIFE OFFICIAL STORE':['MAP0067'],
 'DAIONICA OFFICIAL STORE':['MAP0069'],'DANCOLY PARIS HQ':['MAP0070'],'MIX ENZYME':['MAP0003'],
 'PAW PAWS OFFICIAL':['MAP0075'],'PETAVIT':['MAP0076'],'WINSENSE':['MAP0077'],
 'HAIR FACTORY OFFICIAL':['MAP0071'],'J PACKAGING':['MAP0072'],'JEN MOMMY ESSENTIAL OIL':['MAP0073']},
 'SG': {'BERLANCO SG':['MAP0066'],'BONLIFE SG':['MAP0068'],'DRSMILE SG':['MAP0034'],
 'GOHERB SG':['MAP0001'],'ILADY SG':['MAP0035','MAP0036'],'KATA MARINE':['MAP0037'],
 'MCS SG':['MAP0040'],'MFORMULA SG':['MAP0038'],'MOESIE SG':[],'NINOKO SG':['MAP0045'],
 'SCALE STORY SG':['MAP0041'],'SKINDAE SG':['MAP0042'],'ZEERO SKINCARE':['MAP0043','MAP0044']}
}
rows=[]
excluded=[]
for market in ['MY','SG']:
 tab=f'SHOPEE {market} (JOLIN YONG)'
 for rowno, cells in enumerate(book[tab].iter_rows(values_only=True),1):
  if rowno==1 or not any(v is not None for v in cells): continue
  entity, project, relation, commission, version = cells[:5]
  if project=='(AFFILIATE)':
   excluded.append({'tab':tab,'row':rowno,'project':project,'reason':'Affiliate billing, not a Shopee store'})
   continue
  key=project.upper()
  assert key in crosswalk[market], (market,key)
  ids=crosswalk[market][key]
  candidates=[]
  for mid in ids:
   r=mapping[mid]
   assert r[3]==market and r[4]=='Shopee'
   candidates.append({'mappingId':mid,'username':r[6] or '', 'storeId':None,'storeName':r[5], 'projectId':r[1]})
  if len(ids)>1:
   state='ambiguous'
   reason='Shared store with Placenta and Slimpro; allocation required.' if key=='MIZINO PLACENTA' else 'Multiple store candidates; confirm billing allocation.'
  elif ids:
   state='candidate';reason='Store candidate found; billing identity needs confirmation.'
  else:
   state='missing';reason='Store mapping missing; confirm username and store.'
  rows.append({'id':f'roster-2026-10-{market.lower()}-{rowno}', 'entity':entity,'market':market,
    'project':project,'termsVersion':'NEW' if version=='NEW VER' else 'OLD',
    'mappingStatus':state,'mappingReason':reason,'candidates':candidates,'sourceRow':rowno,
    'internalExternal':relation,'sourceTab':tab})
assert len(rows)==58
snapshot={'tenantId':'j-packaging','invoiceMonth':'2026-10','statementMonth':'2026-09',
 'generatedAt':datetime.now(ZoneInfo('Asia/Kuala_Lumpur')).isoformat(),'rows':rows,
 'sources':[{'file':p.name,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in [billing,master]],
 'excluded':excluded,'mappingPolicy':'Historical candidates only. No financial store binding is confirmed.'}
(root/'app/account-roster-data.json').write_text(json.dumps(snapshot,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'rows':len(rows),'MY':sum(r['market']=='MY' for r in rows),'SG':sum(r['market']=='SG' for r in rows),
 'mapping':{s:sum(r['mappingStatus']==s for r in rows) for s in ['candidate','ambiguous','missing']}}))
